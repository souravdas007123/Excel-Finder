"""Formatted Excel report (Summary / Found / Not Found).

openpyxl ke write-only mode me banta hai: rows seedha stream hoti hain, isliye badi report par bhi RAM kam lagti hai.
Is file me Django import nahi hai.
"""
import io
import os
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="417690")
FOUND_COLUMNS = [("Number", 18), ("Found As", 18), ("File", 38), ("Folder", 55), ("Sheet", 18),
                 ("Row", 8), ("Column", 9), ("File Modified", 20), ("Full Path", 70)]
MISSING_COLUMNS = [("Number", 20)]
FILE_COLUMNS = [("File", 40), ("Folder", 55), ("Numbers Found", 15), ("Total Matches", 15), ("File Modified", 20),
                ("Full Path", 70)]
LINK_FONT = Font(color="0563C1", underline="single")
EXCEL_MAX_DATA_ROWS = 1_000_000


def _cell(ws, value, bold=False, size=11):
    c = WriteOnlyCell(ws, value=value)
    c.font = Font(bold=bold, size=size)
    return c


def _file_link_cell(ws, folder, name):
    """File ke naam par click karke Excel file khul jaye (file:// link). Link na ban sake toh plain text."""
    c = WriteOnlyCell(ws, value=name)
    try:
        c.hyperlink = Path(os.path.join(folder, name)).as_uri()
        c.font = LINK_FONT
    except ValueError:   # relative / ajeeb path
        pass
    return c


def _prepare(ws, columns):
    """Widths, freeze panes aur header row. Ye pehli data row se PEHLE hona zaruri hai (write-only mode)."""
    for i, (_, width) in enumerate(columns, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"
    header = []
    for label, _ in columns:
        c = WriteOnlyCell(ws, value=label)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.alignment = Alignment(vertical="center")
        header.append(c)
    ws.append(header)


def build_report(chunks, searched, skipped, duplicates, max_hits, file_stats=None, last10=None):
    """chunks: iterable of (numbers, matches, totals, files). Returns .xlsx file ke bytes."""
    wb = Workbook(write_only=True)
    ws_sum = wb.create_sheet("Summary")
    ws_files = wb.create_sheet("By File")
    ws_found = wb.create_sheet("Found")
    ws_missing = wb.create_sheet("Not Found")

    ws_sum.column_dimensions["A"].width = 52
    ws_sum.column_dimensions["B"].width = 28
    _prepare(ws_files, FILE_COLUMNS)
    _prepare(ws_found, FOUND_COLUMNS)
    _prepare(ws_missing, MISSING_COLUMNS)

    found = missing = found_rows = capped = 0
    row_limit_hit = False

    for numbers, matches, totals, _files in chunks:
        for n in numbers:
            if totals[n]:
                found += 1
                if totals[n] > max_hits:
                    capped += 1
                for m in matches[n]:
                    if found_rows >= EXCEL_MAX_DATA_ROWS:
                        row_limit_hit = True
                        break
                    ws_found.append([n, m.get("stored", ""), m["file"], m["folder"], m["sheet"], m["row"], m["column"],
                                     m["modified"], os.path.join(m["folder"], m["file"])])
                    found_rows += 1
            else:
                ws_missing.append([n])
                missing += 1

    file_rows = 0   # By File sheet: kis file me kitne numbers mile (chunks khatam hone ke baad poora count pata hota hai)
    for fs in sorted((file_stats or {}).values(), key=lambda f: (-f["numbers_count"], f["file"].lower())):
        ws_files.append([_file_link_cell(ws_files, fs["folder"], fs["file"]), fs["folder"], fs["numbers_count"],
                         fs["matches"], fs["modified"], os.path.join(fs["folder"], fs["file"])])
        file_rows += 1
    ws_files.auto_filter.ref = f"A1:{get_column_letter(len(FILE_COLUMNS))}{file_rows + 1}"
    ws_found.auto_filter.ref = f"A1:{get_column_letter(len(FOUND_COLUMNS))}{found_rows + 1}"
    ws_missing.auto_filter.ref = f"A1:A{missing + 1}"

    ws_sum.append([_cell(ws_sum, "Number Search Report", bold=True, size=14)])
    ws_sum.append([])
    match_mode = "-" if last10 is None else (
        "Last 10 digits (country code / leading 0 ignored)" if last10 else "Exact number")
    stats = [
        ("Generated on", datetime.now().strftime("%d %b %Y, %H:%M")),
        ("Numbers searched", searched),
        ("Match mode", match_mode),
        ("Found", found),
        ("Not found", missing),
        ("Files containing matches", file_rows),
        ("Duplicates removed", duplicates),
        ("Ignored entries (too short)", len(skipped)),
    ]
    if capped:
        stats.append((f"Numbers with more than {max_hits} matches (first {max_hits} listed)", capped))
    if row_limit_hit:
        stats.append(("Note", "Excel row limit reached, list truncated"))
    for label, value in stats:
        ws_sum.append([_cell(ws_sum, label, bold=True), value])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()