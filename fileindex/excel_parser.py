"""
Excel reading logic. IMPORTANT: yahan Django import NAHI karna,
kyunki ye file process-pool workers me alag process me chalti hai.
"""
import re
from importlib import import_module
from datetime import date, datetime, time, timedelta

try:
    from python_calamine import CalamineWorkbook   # Rust based, bahut fast
    HAS_CALAMINE = True
except ImportError:
    CalamineWorkbook = None
    HAS_CALAMINE = False

MAX_DIGITS = 32
_SEPARATORS = re.compile(r"[\s\-+()]")
_DIGITS = re.compile(r"\d+")
_SKIP_TYPES = (datetime, date, time, timedelta, bool)


def extract_numbers(text, min_digits):
    """Text cell se numbers. '98563 25417', '9856325417.0' sab '9856325417' ban jaate hain."""
    text = text.strip()
    if not text:
        return []
    if text.endswith(".0"):
        text = text[:-2]
    compact = _SEPARATORS.sub("", text)
    found = [compact] if (compact.isascii() and compact.isdigit()) else _DIGITS.findall(text)
    return [n for n in found if min_digits <= len(n) <= MAX_DIGITS]


# ---------------------------------------------------------------- readers
def _calamine_rows(path):
    wb = CalamineWorkbook.from_path(path)
    try:
        for name in wb.sheet_names:
            # skip_empty_area=False => row/column numbers Excel ke asli numbers se match karte hain
            rows = wb.get_sheet_by_name(name).to_python(skip_empty_area=False)
            for r, row in enumerate(rows, 1):
                yield name, r, row
            del rows
    finally:
        close = getattr(wb, "close", None)
        if close:
            close()


def _legacy_rows(path):
    if path.lower().endswith(".xls"):
        xlrd = import_module("xlrd")
        book = xlrd.open_workbook(path, on_demand=True)
        try:
            for sheet in book.sheets():
                for r in range(sheet.nrows):
                    yield sheet.name, r + 1, sheet.row_values(r)
        finally:
            book.release_resources()
    else:
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            for ws in wb.worksheets:
                for r, row in enumerate(ws.iter_rows(values_only=True), 1):
                    yield ws.title, r, row
        finally:
            wb.close()


# ---------------------------------------------------------------- parsing
def _collect(rows, min_digits):
    hits = set()   # (number, sheet, row, col) - duplicates automatically hat jaate hain
    for sheet, row_no, values in rows:
        sheet = sheet[:100]
        for col_no, v in enumerate(values, 1):
            t = type(v)
            if t is str:                      # sabse common
                if v:
                    for n in extract_numbers(v, min_digits):
                        hits.add((n, sheet, row_no, col_no))
            elif t is int:                    # fast path: regex ki zarurat nahi
                s = str(abs(v))
                if min_digits <= len(s) <= MAX_DIGITS:
                    hits.add((s, sheet, row_no, col_no))
            elif t is float:
                if v.is_integer():
                    s = str(int(abs(v)))
                    if min_digits <= len(s) <= MAX_DIGITS:
                        hits.add((s, sheet, row_no, col_no))
                else:
                    for n in extract_numbers(str(v), min_digits):
                        hits.add((n, sheet, row_no, col_no))
            # date/bool/None/baaki types skip
    # Sorted: DB me number ke index me ek order me insert hota hai, set ke random order se kaafi tez (badi file ~30%).
    # Ye kaam worker process me hota hai, isliye main process (DB likhne wala) par bojh nahi badhta.
    return sorted(hits)


def parse_file(path, min_digits):
    """Ek Excel file se saare numbers nikalta hai. Process pool worker isi ko chalata hai."""
    readers = []
    if HAS_CALAMINE:
        readers.append(_calamine_rows)
    if not path.lower().endswith(".xlsb"):
        readers.append(_legacy_rows)          # calamine fail ho toh purana reader try karo

    last_exc = None
    for reader in readers:
        try:
            return _collect(reader(path), min_digits)
        except Exception as exc:
            last_exc = exc
    raise last_exc or RuntimeError("No Excel reader available for " + path)


# ---------------------------------------------------------------- single row (View full row)
def _cell_text(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def read_row(path, sheet_name, row_no):
    """Ek row + sheet ki pehli row (header). Click par hi chalta hai, DB me store nahi hota."""
    headers, values = None, None
    if HAS_CALAMINE:
        try:
            wb = CalamineWorkbook.from_path(path)
            rows = wb.get_sheet_by_name(sheet_name).to_python(skip_empty_area=False)
            headers = rows[0] if rows else []
            values = rows[row_no - 1] if 0 < row_no <= len(rows) else []
        except Exception:
            if path.lower().endswith(".xlsb"):
                raise
            headers = None

    if headers is None:
        if path.lower().endswith(".xls"):
            xlrd = import_module("xlrd")
            book = xlrd.open_workbook(path, on_demand=True)
            try:
                sh = book.sheet_by_name(sheet_name)
                headers = sh.row_values(0) if sh.nrows else []
                values = sh.row_values(row_no - 1) if 0 < row_no <= sh.nrows else []
            finally:
                book.release_resources()
        else:
            from openpyxl import load_workbook
            wb = load_workbook(path, read_only=True, data_only=True)
            try:
                ws = wb[sheet_name]
                headers = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
                values = next(ws.iter_rows(min_row=row_no, max_row=row_no, values_only=True), ())
            finally:
                wb.close()
    return [_cell_text(h) for h in headers], [_cell_text(v) for v in values]