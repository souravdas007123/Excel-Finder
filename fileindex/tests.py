"""Automated tests: scan, search, upload, export, scan control, migration, settings.

Chalane ke liye:   python manage.py test
Koi Excel file ya database chhoona nahi padta: har test apni temporary files aur test database banata hai.
"""
import datetime
import io
import os
import subprocess
import sys
import tempfile
import time
from importlib import import_module
from pathlib import Path
from unittest import mock

import openpyxl
from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import connection
from django.test import SimpleTestCase, TestCase

from . import scanner, services
from .excel_parser import extract_numbers, match_key, parse_file
from .models import FileIndex, NumberIndex, ScanFailure, ScanTask
from .scanner import _friendly_reason, run_scan, scan_lock

BASE_DIR = Path(__file__).resolve().parent.parent


# ------------------------------------------------------------------ helpers
def make_xlsx(path, rows, header=("Name", "Phone"), sheet="Data"):
    """rows: list of rows. Ek simple .xlsx banata hai."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet
    if header:
        ws.append(list(header))
    for r in rows:
        ws.append(list(r) if isinstance(r, (list, tuple)) else [r])
    wb.save(path)
    return str(path)


def xlsx_bytes(rows):
    buf = io.BytesIO()
    wb = openpyxl.Workbook()
    for r in rows:
        wb.active.append(r if isinstance(r, (list, tuple)) else [r])
    wb.save(buf)
    return buf.getvalue()


class TempDirMixin:
    def make_tmp(self):
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        return Path(tmp.name)


class LockSafeMixin:
    """Test ke baad scan lock / cancel flag saaf, taaki ek test dusre ko na bigade."""

    def setUp(self):
        super().setUp()
        scanner.cancel_event.clear()
        self.addCleanup(self._reset_scan_state)

    @staticmethod
    def _reset_scan_state():
        scanner.cancel_event.clear()
        if scan_lock.locked():
            try:
                scan_lock.release()
            except RuntimeError:
                pass


def index_file(name, numbers, folder="/data", mtime=None):
    """Seedha database me ek file aur uske numbers daalo (scan ke bina, tez)."""
    f = FileIndex.objects.create(file_name=name, file_path=f"{folder}/{name}", file_size=1,
                                 file_mtime=mtime or time.time())
    for i, n in enumerate(numbers, start=2):
        NumberIndex.objects.create(file=f, number=str(n), match_key=match_key(str(n)), sheet="Data", row=i, col=2)
    return f


# ------------------------------------------------------------------ parsing
class NumberExtractionTests(SimpleTestCase):
    def test_plain_and_formatted_numbers(self):
        self.assertEqual(extract_numbers("9856325417", 5), ["9856325417"])
        self.assertEqual(extract_numbers("98563 25417", 5), ["9856325417"])
        self.assertEqual(extract_numbers("+91 98563-25417", 5), ["919856325417"])
        self.assertEqual(extract_numbers("(098) 563 25417", 5), ["09856325417"])
        self.assertEqual(extract_numbers("9856325417.0", 5), ["9856325417"])

    def test_text_with_several_numbers_and_short_ones_ignored(self):
        self.assertEqual(extract_numbers("call 9856325417 or 8569741235 ext 12", 5), ["9856325417", "8569741235"])
        self.assertEqual(extract_numbers("1234", 5), [])
        self.assertEqual(extract_numbers("   ", 5), [])

    def test_very_long_digit_runs_are_ignored(self):
        self.assertEqual(extract_numbers("1" * 40, 5), [])

    def test_match_key_ignores_country_code_and_leading_zero(self):
        key = "9856325417"
        for stored in ("9856325417", "919856325417", "09856325417", "0919856325417"):
            self.assertEqual(match_key(stored), key)
        self.assertEqual(match_key("12345"), "12345")       # chhota number jaisa hai waisa
        self.assertEqual(match_key("123456789"), "123456789")


class ParseFileTests(TempDirMixin, SimpleTestCase):
    def test_reads_ints_floats_text_and_skips_dates_and_short_numbers(self):
        tmp = self.make_tmp()
        path = make_xlsx(tmp / "a.xlsx", [
            ["Ravi", 9856325417],                      # int
            ["Sita", 8569741235.0],                    # float jo poora number hai
            ["Amit", "+91 98765 43210"],               # text
            ["Date", datetime.date(2026, 1, 1)],       # date skip
            ["Short", 1234],                           # chhota skip
        ])
        hits = parse_file(path, 5)
        numbers = sorted({h[0] for h in hits})
        self.assertEqual(numbers, ["8569741235", "919876543210", "9856325417"])

    def test_row_and_column_numbers_match_excel(self):
        tmp = self.make_tmp()
        path = make_xlsx(tmp / "a.xlsx", [["x", 9856325417]])   # header row 1, data row 2, column B
        (hit,) = parse_file(path, 5)
        self.assertEqual(hit, ("9856325417", "Data", 2, 2))

    def test_result_is_sorted_and_deduplicated(self):
        tmp = self.make_tmp()
        path = make_xlsx(tmp / "a.xlsx", [[9000000002, 9000000001], [9000000001, 9000000002]])
        hits = parse_file(path, 5)
        self.assertEqual(hits, sorted(hits))
        self.assertEqual(len(hits), len(set(hits)))

    def test_multiple_sheets(self):
        tmp = self.make_tmp()
        wb = openpyxl.Workbook()
        wb.active.title = "One"
        wb.active.append([9000000001])
        wb.create_sheet("Two").append([9000000002])
        path = str(tmp / "m.xlsx")
        wb.save(path)
        sheets = {h[0]: h[1] for h in parse_file(path, 5)}
        self.assertEqual(sheets, {"9000000001": "One", "9000000002": "Two"})

    def test_corrupt_file_raises(self):
        tmp = self.make_tmp()
        bad = tmp / "bad.xlsx"
        bad.write_text("not an excel file")
        with self.assertRaises(Exception):
            parse_file(str(bad), 5)


class ParseInputNumbersTests(SimpleTestCase):
    def test_splits_on_comma_semicolon_and_newline(self):
        nums, skipped, dups = services._parse_numbers("9856325417, 8569741235;9000000001\n9000000002")
        self.assertEqual(nums, ["9856325417", "8569741235", "9000000001", "9000000002"])
        self.assertEqual((skipped, dups), ([], 0))

    def test_short_entries_are_skipped_and_reported(self):
        nums, skipped, _ = services._parse_numbers("9856325417\n123\nabc")
        self.assertEqual(nums, ["9856325417"])
        self.assertEqual(skipped, ["123", "abc"])

    def test_exact_mode_keeps_country_code_variants_apart(self):
        nums, _, dups = services._parse_numbers("9856325417\n919856325417\n9856325417", last10=False)
        self.assertEqual(nums, ["9856325417", "919856325417"])
        self.assertEqual(dups, 1)

    def test_last10_mode_treats_variants_as_duplicates(self):
        nums, _, dups = services._parse_numbers("9856325417\n919856325417\n09856325417", last10=True)
        self.assertEqual(nums, ["9856325417"])
        self.assertEqual(dups, 2)

    def test_formatted_numbers_are_cleaned(self):
        nums, _, _ = services._parse_numbers("+91 98563-25417")
        self.assertEqual(nums, ["919856325417"])


class FriendlyReasonTests(SimpleTestCase):
    def test_reasons(self):
        self.assertIn("corrupted", _friendly_reason(Exception("Zip error: invalid Zip archive")))
        self.assertIn("Password", _friendly_reason(Exception("workbook is encrypted")))
        self.assertIn("permission", _friendly_reason(PermissionError("denied")))
        self.assertIn("moved or deleted", _friendly_reason(FileNotFoundError("gone")))
        self.assertIn("too big", _friendly_reason(MemoryError()))
        self.assertTrue(_friendly_reason(ValueError("weird")).startswith("ValueError"))


# ------------------------------------------------------------------ scanning
@mock.patch.object(scanner, "WORKERS", 1)   # tests me process pool nahi (tez aur simple)
class ScanTests(LockSafeMixin, TempDirMixin, TestCase):
    def scan(self, root):
        task = ScanTask.objects.create(status="Running")
        self.assertTrue(scan_lock.acquire(blocking=False))
        run_scan(task.id, str(root))
        task.refresh_from_db()
        return task

    def test_scan_indexes_files_numbers_and_match_keys(self):
        tmp = self.make_tmp()
        make_xlsx(tmp / "a.xlsx", [["x", 9856325417], ["y", 919000000001]])
        (tmp / "sub").mkdir()
        make_xlsx(tmp / "sub" / "b.xlsx", [["z", 8569741235]])
        task = self.scan(tmp)

        self.assertEqual(task.status, "Completed")
        self.assertEqual((task.files_indexed, task.files_skipped, task.files_failed), (2, 0, 0))
        self.assertEqual(FileIndex.objects.count(), 2)
        keys = dict(NumberIndex.objects.values_list("number", "match_key"))
        self.assertEqual(keys, {"9856325417": "9856325417", "919000000001": "9000000001", "8569741235": "8569741235"})
        self.assertFalse(scan_lock.locked(), "scan ke baad lock free hona chahiye")

    def test_temp_lock_files_and_other_extensions_are_ignored(self):
        tmp = self.make_tmp()
        make_xlsx(tmp / "real.xlsx", [["x", 9856325417]])
        make_xlsx(tmp / "~$real.xlsx", [["x", 9111111111]])      # Excel ki temporary lock file
        (tmp / "notes.txt").write_text("9222222222")
        task = self.scan(tmp)
        self.assertEqual(task.files_indexed, 1)
        self.assertEqual(list(FileIndex.objects.values_list("file_name", flat=True)), ["real.xlsx"])

    def test_rescan_skips_unchanged_reindexes_changed_and_removes_deleted(self):
        tmp = self.make_tmp()
        a = make_xlsx(tmp / "a.xlsx", [["x", 9856325417]])
        b = make_xlsx(tmp / "b.xlsx", [["x", 8569741235]])
        self.scan(tmp)

        task = self.scan(tmp)                                    # kuch nahi badla
        self.assertEqual((task.files_indexed, task.files_skipped), (0, 2))

        make_xlsx(a, [["x", 9000000009]])                        # a badli
        os.utime(a, (time.time() + 5, time.time() + 5))
        os.remove(b)                                             # b delete
        task = self.scan(tmp)
        self.assertEqual((task.files_indexed, task.files_skipped), (1, 0))
        self.assertEqual(set(NumberIndex.objects.values_list("number", flat=True)), {"9000000009"})
        self.assertEqual(list(FileIndex.objects.values_list("file_name", flat=True)), ["a.xlsx"])

    def test_failed_files_are_listed_with_reason_and_do_not_stop_the_scan(self):
        tmp = self.make_tmp()
        make_xlsx(tmp / "good.xlsx", [["x", 9856325417]])
        (tmp / "broken.xlsx").write_text("this is not a spreadsheet")
        (tmp / "empty.xlsx").write_bytes(b"")
        task = self.scan(tmp)

        self.assertEqual(task.status, "Completed")
        self.assertEqual((task.files_indexed, task.files_failed), (1, 2))
        failures = {os.path.basename(f.file_path): f.reason for f in ScanFailure.objects.filter(task=task)}
        self.assertEqual(set(failures), {"broken.xlsx", "empty.xlsx"})
        self.assertTrue(all(r for r in failures.values()))
        self.assertIn("corrupted", failures["broken.xlsx"])

    def test_missing_folder_marks_task_as_error_and_releases_lock(self):
        task = self.scan(self.make_tmp() / "does-not-exist")
        self.assertEqual(task.status, "Error")
        self.assertIn("Folder not found", task.message)
        self.assertFalse(scan_lock.locked())

    def test_stop_keeps_finished_files_and_never_leaves_a_half_indexed_file(self):
        tmp = self.make_tmp()
        for i in range(4):
            make_xlsx(tmp / f"f{i}.xlsx", [["x", 9000000000 + i * 10 + k] for k in range(5)])
        real = scanner.parse_file
        calls = []

        def stop_after_two(path, min_digits):
            calls.append(path)
            result = real(path, min_digits)
            if len(calls) == 2:
                scanner.cancel_event.set()         # user ne Stop dabaya
            return result

        with mock.patch.object(scanner, "parse_file", stop_after_two):
            task = self.scan(tmp)

        self.assertEqual(task.status, "Cancelled")
        self.assertLess(len(calls), 4, "Stop ke baad baaki files nahi padhni chahiye")
        for f in FileIndex.objects.all():          # jo index hui wo poori ho
            self.assertEqual(f.numbers.count(), 5)
        self.assertFalse(scan_lock.locked())

    def test_progress_total_is_counted_for_the_time_estimate(self):
        tmp = self.make_tmp()
        for i in range(3):
            make_xlsx(tmp / f"f{i}.xlsx", [["x", 9856325400 + i]])
        task = self.scan(tmp)
        for _ in range(40):                         # ginti wala thread kuch millisecond me khatam hota hai
            if scanner.scan_progress["total"] is not None:
                break
            time.sleep(0.05)
        self.assertEqual(scanner.scan_progress["total"], 3)
        self.assertEqual(scanner.scan_progress["task_id"], task.id)

    def test_new_scan_resets_progress(self):
        tmp = self.make_tmp()
        make_xlsx(tmp / "a.xlsx", [["x", 9856325417]])
        self.scan(tmp)
        first = scanner.scan_progress["run"]
        self.scan(tmp)
        self.assertIsNot(scanner.scan_progress["run"], first)


# ------------------------------------------------------------------ search API
class ServiceTestCase(LockSafeMixin, TempDirMixin, TestCase):
    @classmethod
    def setUpTestData(cls):
        User = get_user_model()
        cls.admin = User.objects.create_superuser("admin", "a@example.com", "pw")
        cls.staff = User.objects.create_user("staff", "s@example.com", "pw", is_staff=True)


class SearchTests(ServiceTestCase):
    def setUp(self):
        super().setUp()
        self.alpha = index_file("alpha.xlsx", [9000000001, 9000000002, 9000000003])
        self.beta = index_file("beta.xlsx", [9000000002, 919856325417])
        self.gamma = index_file("gamma.xlsx", ["09111111111"])

    def search(self, numbers, last10=True):
        return services.bulk_search_payload(numbers, last10)

    def test_counts_files_and_not_found(self):
        j = self.search("9000000001\n9000000002\n9000000003\n9999999999")
        s = j["summary"]
        self.assertEqual((s["searched"], s["found"], s["not_found"], s["files"]), (4, 3, 1, 2))
        self.assertEqual(j["not_found"], ["9999999999"])
        files = {f["file"]: f for f in j["files"]}
        self.assertEqual(files["alpha.xlsx"]["numbers_count"], 3)
        self.assertEqual(files["beta.xlsx"]["numbers_count"], 1)
        self.assertEqual([f["file"] for f in j["files"]], ["alpha.xlsx", "beta.xlsx"])   # zyada numbers pehle

    def test_per_number_details_and_multi_file_counts(self):
        j = self.search("9000000002")
        (r,) = j["results"]
        self.assertTrue(r["found"])
        self.assertEqual((r["total"], r["file_count"]), (2, 2))
        self.assertEqual({m["file"] for m in r["matches"]}, {"alpha.xlsx", "beta.xlsx"})
        m = r["matches"][0]
        self.assertEqual((m["sheet"], m["column"]), ("Data", "B"))

    def test_last10_matches_country_code_and_leading_zero(self):
        j = self.search("9856325417\n9111111111")
        self.assertEqual(j["summary"]["found"], 2)
        stored = {m["stored"] for r in j["results"] for m in r["matches"]}
        self.assertEqual(stored, {"919856325417", "09111111111"})

    def test_exact_mode_does_not_match_variants(self):
        j = self.search("9856325417\n9111111111", last10=False)
        self.assertEqual((j["summary"]["found"], j["summary"]["last10"]), (0, False))
        j = self.search("919856325417", last10=False)
        self.assertEqual(j["summary"]["found"], 1)
        self.assertEqual(j["results"][0]["matches"][0]["stored"], "")     # jaisa likha waisa mila

    def test_duplicates_ignored_short_numbers_reported_and_limit(self):
        j = self.search("9856325417, 919856325417, 12, 9856325417")
        self.assertEqual(j["summary"]["searched"], 1)
        self.assertEqual(j["summary"]["duplicates"], 2)
        self.assertEqual(j["summary"]["skipped"], ["12"])
        with mock.patch.object(services, "MAX_NUMBERS", 3):
            j = self.search("\n".join(str(9000000100 + i) for i in range(10)))
        self.assertTrue(j["summary"]["truncated"])
        self.assertEqual(j["summary"]["searched"], 3)

    def test_empty_input_is_harmless(self):
        j = self.search("")
        self.assertEqual((j["summary"]["searched"], j["files"], j["results"]), (0, [], []))

    def test_numbers_are_chunked_without_losing_matches(self):
        with mock.patch.object(services, "CHUNK", 2):
            j = self.search("9000000001\n9000000002\n9000000003\n9856325417\n9111111111")
        self.assertEqual(j["summary"]["found"], 5)

    def test_file_locations(self):
        items = {i["number"]: i["locations"] for i in services.file_locations(self.beta.id, "9000000002\n9856325417")}
        self.assertEqual(set(items), {"9000000002", "9856325417"})
        self.assertEqual(items["9856325417"][0]["stored"], "919856325417")
        self.assertEqual(items["9000000002"][0]["stored"], "")

    def test_row_cells_reads_the_real_file(self):
        tmp = self.make_tmp()
        path = make_xlsx(tmp / "real.xlsx", [["Ravi", 9856325417]])
        f = FileIndex.objects.create(file_name="real.xlsx", file_path=path)
        _file, cells, error, status = services.row_cells(f.id, "Data", 2, 2)
        self.assertEqual((error, status), (None, None))
        self.assertEqual([c["value"] for c in cells], ["Ravi", "9856325417"])
        self.assertEqual([c["header"] for c in cells], ["Name", "Phone"])
        self.assertEqual([c["hit"] for c in cells], [False, True])

    def test_row_cells_errors(self):
        self.assertEqual(services.row_cells(99999, "Data", 2)[2:], ("File not in index", 404))
        gone = FileIndex.objects.create(file_name="gone.xlsx", file_path="/no/such/gone.xlsx")
        self.assertEqual(services.row_cells(gone.id, "Data", 2)[3], 404)


class ExportTests(ServiceTestCase):
    def setUp(self):
        super().setUp()
        tmp = self.make_tmp()
        self.folder = str(tmp)
        index_file("alpha book.xlsx", [9856325417], folder=self.folder)
        index_file("with91.xlsx", [919856325417], folder=self.folder)

    def load(self, data):
        return openpyxl.load_workbook(io.BytesIO(data))

    def test_export_has_all_sheets_links_paths_and_match_info(self):
        wb = self.load(services.export_report("9856325417\n1112223334"))
        self.assertEqual(wb.sheetnames, ["Summary", "By File", "Found", "Not Found"])

        by_file = list(wb["By File"].iter_rows(min_row=1, values_only=False))
        self.assertEqual([c.value for c in by_file[0]][-1], "Full Path")
        names = {r[0].value: r for r in by_file[1:]}
        self.assertEqual(set(names), {"alpha book.xlsx", "with91.xlsx"})
        link = names["alpha book.xlsx"][0].hyperlink.target
        self.assertTrue(link.startswith("file:///") and link.endswith("alpha%20book.xlsx"), link)
        self.assertEqual(names["alpha book.xlsx"][-1].value, os.path.join(self.folder, "alpha book.xlsx"))

        found = list(wb["Found"].iter_rows(min_row=2, values_only=True))
        self.assertEqual(sorted(r[1] or "" for r in found), ["", "919856325417"])      # Found As

        self.assertEqual([r[0] for r in wb["Not Found"].iter_rows(min_row=2, values_only=True)], ["1112223334"])
        summary = {r[0]: r[1] for r in wb["Summary"].iter_rows(min_row=3, values_only=True) if r[0]}
        self.assertEqual((summary["Found"], summary["Not found"]), (1, 1))
        self.assertTrue(summary["Match mode"].startswith("Last 10 digits"))

    def test_export_exact_mode(self):
        wb = self.load(services.export_report("9856325417", last10=False))
        found = list(wb["Found"].iter_rows(min_row=2, values_only=True))
        self.assertEqual([r[2] for r in found], ["alpha book.xlsx"])
        summary = {r[0]: r[1] for r in wb["Summary"].iter_rows(min_row=3, values_only=True) if r[0]}
        self.assertEqual(summary["Match mode"], "Exact number")

    def test_file_report_from_uploaded_excel(self):
        up = SimpleUploadedFile("nums.xlsx", xlsx_bytes([[9856325417], [1112223334]]))
        data, error, status = services.file_report(up)
        self.assertEqual((error, status), (None, None))
        wb = self.load(data)
        self.assertEqual(len(list(wb["Found"].iter_rows(min_row=2))), 2)
        self.assertEqual(len(list(wb["Not Found"].iter_rows(min_row=2))), 1)

    def test_file_report_with_no_numbers_is_an_error(self):
        data, error, status = services.file_report(SimpleUploadedFile("nums.txt", b"no digits here"))
        self.assertEqual((data, status), (None, 400))
        self.assertIn("No numbers", error)

    def test_report_filename_is_dated(self):
        self.assertRegex(services.xlsx_filename(), r"^number_search_report_\d{8}_\d{4}\.xlsx$")


class UploadTests(ServiceTestCase):
    def extract(self, name, content, first_col=False, last10=True):
        return services.read_upload_numbers(SimpleUploadedFile(name, content), first_col, last10)

    def numbers(self, *args, **kw):
        result, error, _status = self.extract(*args, **kw)
        self.assertIsNone(error)
        return result[1]

    def test_csv_with_quotes_and_header(self):
        self.assertEqual(self.numbers("n.csv", b'phone,name\n"9856325417",Ravi\n8569741235,Sita\n'), ["9856325417", "8569741235"])

    def test_txt_dedupes_and_keeps_order(self):
        result, _, _ = self.extract("n.txt", b"9000000020\n9000000020\n8000000006\n")
        self.assertEqual(result, ("n.txt", ["9000000020", "8000000006"], 1))

    def test_first_column_only(self):
        self.assertEqual(self.numbers("n.csv", b"9856325417,8569741235\n9000000001,9000000002\n", first_col=True),
                         ["9856325417", "9000000001"])
        self.assertEqual(self.numbers("n.xlsx", xlsx_bytes([[9856325417, 8569741235]]), first_col=True), ["9856325417"])

    def test_excel_upload(self):
        self.assertEqual(self.numbers("n.xlsx", xlsx_bytes([[9856325417], ["+91 98765 43210"]])), ["9856325417", "919876543210"])

    def test_country_code_variants_are_merged_by_default(self):
        self.assertEqual(self.numbers("n.txt", b"9856325417\n919856325417\n"), ["9856325417"])
        self.assertEqual(len(self.numbers("n.txt", b"9856325417\n919856325417\n", last10=False)), 2)

    def test_bad_uploads(self):
        self.assertEqual(self.extract("n.pdf", b"x")[1:], ("Unsupported file. Use .xlsx, .xls, .csv or .txt", 400))
        self.assertEqual(services.read_upload_numbers(None)[1:], ("No file received", 400))
        self.assertEqual(self.extract("bad.xlsx", b"not excel")[2], 400)
        with mock.patch.object(services, "MAX_UPLOAD_MB", 0.0001):
            self.assertEqual(self.extract("n.txt", b"9" * 500)[2], 413)


class ScanControlTests(ServiceTestCase):
    def test_clear_index_removes_everything(self):
        index_file("a.xlsx", [9856325417, 9000000001])
        index_file("b.xlsx", [8569741235])
        removed, error, _ = services.clear_index(self.admin)
        self.assertEqual((removed, error), (2, None))
        self.assertEqual((FileIndex.objects.count(), NumberIndex.objects.count()), (0, 0))
        self.assertFalse(scan_lock.locked())

    def test_clear_index_refuses_while_a_scan_runs(self):
        index_file("a.xlsx", [9856325417])
        scan_lock.acquire()
        self.assertEqual(services.clear_index(self.admin)[1:], ("A scan is running. Please wait for it to finish.", 409))
        self.assertEqual(FileIndex.objects.count(), 1)

    def test_clear_index_needs_delete_permission(self):
        index_file("a.xlsx", [9856325417])
        self.assertEqual(services.clear_index(self.staff)[2], 403)       # staff hai par delete permission nahi
        self.assertEqual(FileIndex.objects.count(), 1)

    def test_stop_scan(self):
        self.assertEqual(services.stop_scan(), (False, "No scan is running"))      # kuch chal hi nahi raha
        scan_lock.acquire()
        self.assertEqual(services.stop_scan(), (True, None))
        self.assertTrue(scanner.cancel_event.is_set())

    def test_start_scan_validation(self):
        self.assertEqual(services.start_scan(self.admin, "nope")[1:], ("Invalid location", 400))
        self.assertEqual(services.start_scan(self.admin, "__custom__", "/definitely/not/here")[2], 400)
        self.assertEqual(services.start_scan(self.staff, "__custom__", "/")[2], 403)      # custom folder sirf superuser

    def test_start_scan_refuses_second_scan(self):
        scan_lock.acquire()
        self.assertEqual(services.start_scan(self.admin, "__custom__", str(self.make_tmp()))[2], 409)

    def test_start_scan_starts_a_task(self):
        tmp = self.make_tmp()
        started = []

        def fake_scan(task_id, root):
            started.append((task_id, root))
            scan_lock.release()                   # asli run_scan bhi yahi karta hai

        with mock.patch.object(services, "run_scan", fake_scan):
            ScanTask.objects.create(status="Running")                    # purana atka hua task
            task_id, error, _ = services.start_scan(self.admin, "__custom__", str(tmp))
            self.assertEqual(error, None)
            for _ in range(40):
                if started:
                    break
                time.sleep(0.05)
        self.assertEqual(started[0], (task_id, str(tmp)))
        self.assertEqual(ScanTask.objects.filter(status="Error", message="Interrupted").count(), 1)

    def test_scan_status(self):
        t = ScanTask.objects.create(status="Completed", files_indexed=3, files_failed=1)
        j = services.scan_status(t)
        self.assertEqual((j["status"], j["files_indexed"], j["files_failed"]), ("Completed", 3, 1))
        self.assertIsNone(j["files_total"])

    def test_scan_status_reports_total_for_the_running_scan(self):
        t = ScanTask.objects.create(status="Running")
        with mock.patch.dict(scanner.scan_progress, {"task_id": t.id, "total": 120, "counted": 120}):
            j = services.scan_status(t)
        self.assertEqual((j["files_total"], j["files_counted"]), (120, 120))

    def test_browse_folders_lists_sub_folders(self):
        tmp = self.make_tmp()
        (tmp / "Beta").mkdir()
        (tmp / "alpha").mkdir()
        (tmp / ".hidden").mkdir()
        (tmp / "file.txt").write_text("x")
        data, error, _ = services.browse_folders(str(tmp))
        self.assertEqual([f["name"] for f in data["folders"]], ["alpha", "Beta"])     # hidden aur files nahi
        self.assertEqual(services.browse_folders("/no/such")[1:], ("Folder not found", 400))


# ------------------------------------------------------------------ migration + settings
class MigrationTests(TestCase):
    def test_match_key_fill_handles_all_lengths(self):
        migration = import_module("fileindex.migrations.0007_numberindex_match_key")

        class Editor:   # schema_editor ki jagah (asli editor test transaction me chal nahi sakta)
            connection = connection

            @staticmethod
            def execute(sql):
                with connection.cursor() as cur:
                    cur.execute(sql)

        f = FileIndex.objects.create(file_name="a", file_path="/a")
        wanted = {"9856325417": "9856325417", "919856325417": "9856325417", "09856325417": "9856325417",
                  "12345": "12345", "123456789": "123456789", "12345678901234567890": "1234567890"}
        for n in wanted:
            NumberIndex.objects.create(file=f, number=n, match_key="", sheet="S", row=1, col=1)
        migration.fill_match_key(None, Editor)
        self.assertEqual(dict(NumberIndex.objects.values_list("number", "match_key")), wanted)

    def test_fill_agrees_with_python_match_key(self):
        for n in ("9856325417", "919856325417", "12345", "12345678901234567890"):
            self.assertEqual(match_key(n), n[-10:] if len(n) > 10 else n)


class SettingsTests(SimpleTestCase):
    def run_settings(self, **env):
        full = {**os.environ, "DJANGO_SETTINGS_MODULE": "search.settings", "PYTHONPATH": str(BASE_DIR)}
        for k in ("DJANGO_DEBUG", "DJANGO_SECRET_KEY", "DJANGO_ALLOWED_HOSTS"):
            full.pop(k, None)
        full.update(env)
        code = ("from django.conf import settings; "
                "print(settings.DEBUG, settings.ALLOWED_HOSTS, settings.SESSION_COOKIE_SECURE)")
        return subprocess.run([sys.executable, "-c", code], env=full, cwd=BASE_DIR, capture_output=True, text=True)

    def test_local_defaults_unchanged(self):
        r = self.run_settings()
        self.assertEqual(r.returncode, 0, r.stderr)
        self.assertEqual(r.stdout.split()[0], "True")

    def test_production_mode_refuses_the_development_secret_key(self):
        r = self.run_settings(DJANGO_DEBUG="0")
        self.assertNotEqual(r.returncode, 0)
        self.assertIn("DJANGO_SECRET_KEY", r.stderr)

    def test_production_mode_with_key_and_hosts(self):
        r = self.run_settings(DJANGO_DEBUG="0", DJANGO_SECRET_KEY="x" * 60, DJANGO_ALLOWED_HOSTS="a.com, b.com")
        self.assertEqual(r.returncode, 0, r.stderr)
        self.assertIn("['a.com', 'b.com']", r.stdout)
        self.assertTrue(r.stdout.strip().endswith("True"))              # https-only cookies
